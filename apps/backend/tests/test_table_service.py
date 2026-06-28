from datetime import datetime
from io import BytesIO
from typing import cast

import pytest
from openpyxl import Workbook
from sqlalchemy.orm import Session

from featuremaker.models.table import SourceType, StorageType, TableAsset
from featuremaker.schemas.tables import TableAssetDeleteRequest, TableAssetDetail, TableAssetUpdateRequest
from featuremaker.services.table_service import (
    TableAssetImportError,
    TableAssetNameExistsError,
    TableAssetNotFoundError,
    TableAssetService,
)
from featuremaker.storage.local_table_storage import LocalTableStorage
from featuremaker.table_engines.pandas_table_engine import PandasTableEngine


class FakeSession:
    def __init__(self) -> None:
        self.committed = False
        self.rolled_back = False

    def commit(self) -> None:
        self.committed = True

    def rollback(self) -> None:
        self.rolled_back = True


class FakeRepository:
    name_exists = False

    def __init__(self, session: object) -> None:
        self.session = session
        self.created_table_asset: TableAsset | None = None
        self.deleted_table_asset: TableAsset | None = None
        self.updated_table_asset: TableAsset | None = None

    def exists_by_name(self, name: str, *, exclude_id: int | None = None) -> bool:
        return self.name_exists

    def create(self, table_asset: TableAsset) -> TableAsset:
        table_asset.id = 1
        table_asset.source_type = SourceType.UPLOAD_FILE
        table_asset.storage_type = StorageType.LOCAL_CSV
        table_asset.created_at = datetime(2026, 6, 27)
        self.created_table_asset = table_asset
        return table_asset

    def get_by_id(self, table_asset_id: int) -> TableAsset | None:
        if table_asset_id == 404:
            return None
        return TableAsset(
            id=table_asset_id,
            name="customers",
            description="客户表",
            source_type=SourceType.UPLOAD_FILE,
            storage_type=StorageType.LOCAL_CSV,
            storage_uri="data/table-assets/customers.csv",
            schema_json={"columns": []},
            row_count=1,
            created_at=datetime(2026, 6, 27),
        )

    def soft_delete(self, table_asset: TableAsset, *, deleted_by: int | None = None) -> TableAsset:
        table_asset.deleted_at = datetime(2026, 6, 28)
        table_asset.deleted_by = deleted_by
        self.deleted_table_asset = table_asset
        return table_asset

    def update(self, table_asset: TableAsset) -> TableAsset:
        self.updated_table_asset = table_asset
        return table_asset


class FakeStorage:
    def __init__(self) -> None:
        self.deleted_storage_uri: str | None = None
        self.saved_text_content: str | None = None

    def save(self, *, file_name: str, file_obj: object) -> str:
        return "data/table-assets/customers.csv"

    def save_text(self, *, file_name: str, content: str) -> str:
        self.saved_text_content = content
        return "data/table-assets/customers_from_xlsx.csv"

    def delete(self, storage_uri: str) -> None:
        self.deleted_storage_uri = storage_uri


class FakeTableEngine:
    def inspect_csv(self, storage_uri: str) -> tuple[dict, int]:
        return {"columns": [{"name": "customer_id", "type": "integer", "nullable": True}]}, 2


def test_table_asset_service_import_csv_creates_table_asset(monkeypatch):
    fake_session = FakeSession()
    storage = FakeStorage()
    table_engine = FakeTableEngine()

    monkeypatch.setattr(
        "featuremaker.services.table_service.TableAssetRepository",
        FakeRepository,
    )

    service = TableAssetService(
        session=cast(Session, fake_session),
        storage=cast(LocalTableStorage, storage),
        table_engine=cast(PandasTableEngine, table_engine),
    )

    table_asset = service.import_csv(
        file_name="customers.csv",
        file_obj=BytesIO(b"customer_id,name\n1,Alice\n"),
        name="customers",
        description="客户表",
    )

    assert isinstance(table_asset, TableAssetDetail)
    assert table_asset.id == 1
    assert table_asset.name == "customers"
    assert table_asset.row_count == 2
    assert table_asset.table_schema.columns[0].name == "customer_id"
    assert fake_session.committed is True
    assert fake_session.rolled_back is False
    assert storage.deleted_storage_uri is None


def test_table_asset_service_import_xlsx_converts_to_csv_and_creates_table_asset(monkeypatch):
    fake_session = FakeSession()
    storage = FakeStorage()
    table_engine = FakeTableEngine()

    monkeypatch.setattr(
        "featuremaker.services.table_service.TableAssetRepository",
        FakeRepository,
    )

    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["customer_id", "name"])
    worksheet.append([None, None])
    worksheet.append([1, "Alice"])

    xlsx_file = BytesIO()
    workbook.save(xlsx_file)
    workbook.close()
    xlsx_file.seek(0)

    service = TableAssetService(
        session=cast(Session, fake_session),
        storage=cast(LocalTableStorage, storage),
        table_engine=cast(PandasTableEngine, table_engine),
    )

    table_asset = service.import_xlsx(
        file_name="customers.xlsx",
        file_obj=xlsx_file,
        name="customers_from_xlsx",
        description="客户表",
    )

    assert isinstance(table_asset, TableAssetDetail)
    assert table_asset.id == 1
    assert table_asset.name == "customers_from_xlsx"
    assert table_asset.row_count == 2
    assert storage.saved_text_content == "customer_id,name\r\n1,Alice\r\n"
    assert fake_session.committed is True
    assert fake_session.rolled_back is False
    assert storage.deleted_storage_uri is None


def test_table_asset_service_import_xlsx_rejects_empty_header(monkeypatch):
    fake_session = FakeSession()
    storage = FakeStorage()
    table_engine = FakeTableEngine()

    monkeypatch.setattr(
        "featuremaker.services.table_service.TableAssetRepository",
        FakeRepository,
    )

    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append([None, None])

    xlsx_file = BytesIO()
    workbook.save(xlsx_file)
    workbook.close()
    xlsx_file.seek(0)

    service = TableAssetService(
        session=cast(Session, fake_session),
        storage=cast(LocalTableStorage, storage),
        table_engine=cast(PandasTableEngine, table_engine),
    )

    with pytest.raises(TableAssetImportError):
        service.import_xlsx(
            file_name="customers.xlsx",
            file_obj=xlsx_file,
            name="customers_from_xlsx",
            description="客户表",
        )

    assert fake_session.committed is False
    assert fake_session.rolled_back is True
    assert storage.deleted_storage_uri is None


def test_table_asset_service_delete_table_soft_deletes_table_asset(monkeypatch):
    fake_session = FakeSession()
    storage = FakeStorage()
    table_engine = FakeTableEngine()

    monkeypatch.setattr(
        "featuremaker.services.table_service.TableAssetRepository",
        FakeRepository,
    )

    service = TableAssetService(
        session=cast(Session, fake_session),
        storage=cast(LocalTableStorage, storage),
        table_engine=cast(PandasTableEngine, table_engine),
    )

    result = service.delete_table(TableAssetDeleteRequest(id=1))

    assert result.id == 1
    assert result.name == "customers"
    assert fake_session.committed is True
    assert fake_session.rolled_back is False
    assert storage.deleted_storage_uri is None


def test_table_asset_service_delete_table_rejects_missing_table_asset(monkeypatch):
    fake_session = FakeSession()
    storage = FakeStorage()
    table_engine = FakeTableEngine()

    monkeypatch.setattr(
        "featuremaker.services.table_service.TableAssetRepository",
        FakeRepository,
    )

    service = TableAssetService(
        session=cast(Session, fake_session),
        storage=cast(LocalTableStorage, storage),
        table_engine=cast(PandasTableEngine, table_engine),
    )

    with pytest.raises(TableAssetNotFoundError):
        service.delete_table(TableAssetDeleteRequest(id=404))

    assert fake_session.committed is False
    assert fake_session.rolled_back is False


def test_table_asset_service_update_table_updates_metadata(monkeypatch):
    fake_session = FakeSession()
    storage = FakeStorage()
    table_engine = FakeTableEngine()
    FakeRepository.name_exists = False

    monkeypatch.setattr(
        "featuremaker.services.table_service.TableAssetRepository",
        FakeRepository,
    )

    service = TableAssetService(
        session=cast(Session, fake_session),
        storage=cast(LocalTableStorage, storage),
        table_engine=cast(PandasTableEngine, table_engine),
    )

    table_asset = service.update_table(
        TableAssetUpdateRequest(id=1, name="customers_v2", description=None),
    )

    assert table_asset.id == 1
    assert table_asset.name == "customers_v2"
    assert table_asset.description is None
    assert fake_session.committed is True
    assert fake_session.rolled_back is False


def test_table_asset_service_update_table_rejects_duplicate_name(monkeypatch):
    fake_session = FakeSession()
    storage = FakeStorage()
    table_engine = FakeTableEngine()
    FakeRepository.name_exists = True

    monkeypatch.setattr(
        "featuremaker.services.table_service.TableAssetRepository",
        FakeRepository,
    )

    service = TableAssetService(
        session=cast(Session, fake_session),
        storage=cast(LocalTableStorage, storage),
        table_engine=cast(PandasTableEngine, table_engine),
    )

    with pytest.raises(TableAssetNameExistsError):
        service.update_table(TableAssetUpdateRequest(id=1, name="customers_v2"))

    assert fake_session.committed is False
    assert fake_session.rolled_back is False
    FakeRepository.name_exists = False
