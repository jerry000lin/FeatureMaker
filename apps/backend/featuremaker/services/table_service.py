import csv
from io import StringIO
from typing import BinaryIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from featuremaker.models.table import SourceType, StorageType, TableAsset
from featuremaker.repositories.table_asset_repository import TableAssetRepository
from featuremaker.schemas.common import PageParams, PageResponse
from featuremaker.schemas.tables import TableAssetDetail, TableAssetSummary, TablePreviewResponse, TableSchema
from featuremaker.storage.local_table_storage import LocalTableStorage
from featuremaker.table_engines.pandas_table_engine import PandasTableEngine


class TableAssetNotFoundError(Exception):
    """
    表资产不存在。
    """


class TableAssetInvalidFileTypeError(Exception):
    """
    表资产导入文件类型不合法。
    """


class TableAssetNameExistsError(Exception):
    """
    表资产名称已存在。
    """


class TableAssetImportError(Exception):
    """
    表资产导入失败。
    """


def _validate_file_suffix(file_name: str, allowed_suffixes: tuple[str, ...], message: str) -> None:
    if not file_name.lower().endswith(allowed_suffixes):
        raise TableAssetInvalidFileTypeError(message)


class TableAssetService:
    """
    表资产业务服务。
    """

    def __init__(
        self,
        session: Session,
        storage: LocalTableStorage,
        table_engine: PandasTableEngine,
    ) -> None:
        self.session = session
        self.repository = TableAssetRepository(session)
        self.storage = storage
        self.table_engine = table_engine

    def list_tables(self, page: PageParams) -> PageResponse[TableAssetSummary]:
        """
        查询表资产列表。
        """
        table_assets = self.repository.list(offset=page.offset, limit=page.limit)
        total = self.repository.count()

        return PageResponse[TableAssetSummary](
            items=[TableAssetSummary.model_validate(table_asset) for table_asset in table_assets],
            total=total,
            page=page.page,
            page_size=page.page_size,
        )

    def get_table(self, table_asset_id: int) -> TableAssetDetail:
        """
        查询表资产详情。
        """
        table_asset = self.repository.get_by_id(table_asset_id)
        if table_asset is None:
            raise TableAssetNotFoundError(f"表资产不存在: {table_asset_id}")
        return TableAssetDetail.model_validate(table_asset)

    def import_csv(
        self,
        *,
        file_name: str,
        file_obj: BinaryIO,
        name: str,
        description: str | None,
    ) -> TableAssetDetail:
        """
        导入 CSV 文件并创建表资产。
        """
        _validate_file_suffix(file_name, (".csv",), "只支持上传 CSV 文件")
        if self.repository.exists_by_name(name):
            raise TableAssetNameExistsError("表资产名称已存在")

        storage_uri: str | None = None
        try:
            storage_uri = self.storage.save(file_name=file_name, file_obj=file_obj)
            schema_json, row_count = self.table_engine.inspect_csv(storage_uri)
            table_asset = self._create_upload_table_asset(
                name=name,
                description=description,
                storage_uri=storage_uri,
                schema_json=schema_json,
                row_count=row_count,
            )
            self.repository.create(table_asset)
            table_asset_detail = TableAssetDetail.model_validate(table_asset)
            self.session.commit()
            return table_asset_detail
        except Exception as error:
            self.session.rollback()
            if storage_uri is not None:
                self.storage.delete(storage_uri)
            raise TableAssetImportError("CSV 文件导入失败") from error

    def import_xlsx(
        self,
        *,
        file_name: str,
        file_obj: BinaryIO,
        name: str,
        description: str | None,
    ) -> TableAssetDetail:
        """
        导入 XLSX 文件并创建表资产。
        """
        _validate_file_suffix(file_name, (".xlsx",), "只支持上传 XLSX 文件")
        if self.repository.exists_by_name(name):
            raise TableAssetNameExistsError("表资产名称已存在")

        storage_uri: str | None = None
        try:
            csv_content = self._convert_xlsx_to_csv(file_obj)
            storage_uri = self.storage.save_text(file_name=f"{file_name}.csv", content=csv_content)
            schema_json, row_count = self.table_engine.inspect_csv(storage_uri)
            table_asset = self._create_upload_table_asset(
                name=name,
                description=description,
                storage_uri=storage_uri,
                schema_json=schema_json,
                row_count=row_count,
            )
            self.repository.create(table_asset)
            table_asset_detail = TableAssetDetail.model_validate(table_asset)
            self.session.commit()
            return table_asset_detail
        except Exception as error:
            self.session.rollback()
            if storage_uri is not None:
                self.storage.delete(storage_uri)
            raise TableAssetImportError("XLSX 文件导入失败") from error

    def preview_table(self, table_asset_id: int, limit: int) -> TablePreviewResponse:
        """
        查询表资产预览数据。
        """
        table_asset = self.repository.get_by_id(table_asset_id)
        if table_asset is None:
            raise TableAssetNotFoundError(f"表资产不存在: {table_asset_id}")

        rows = self.table_engine.preview_csv(
            table_asset.storage_uri,
            limit=limit,
            schema_json=table_asset.schema_json,
        )
        return TablePreviewResponse(
            id=table_asset.id,
            table_schema=TableSchema.model_validate(table_asset.schema_json),
            rows=rows,
            row_count=table_asset.row_count or 0,
        )

    def _convert_xlsx_to_csv(self, file_obj: BinaryIO) -> str:
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            if worksheet is None:
                raise ValueError("XLSX 文件没有工作表")

            row_iterator = worksheet.iter_rows(values_only=True)
            first_row = next(row_iterator, None)
            if first_row is None:
                raise ValueError("XLSX 文件没有可导入数据")

            header = self._normalize_xlsx_row(first_row)
            if not any(header):
                raise ValueError("XLSX 文件首行表头不能为空")

            output = StringIO()
            writer = csv.writer(output)
            writer.writerow(header)
            for row in row_iterator:
                normalized_row = self._normalize_xlsx_row(row)
                if any(normalized_row):
                    writer.writerow(normalized_row)
            return output.getvalue()
        finally:
            workbook.close()

    def _create_upload_table_asset(
        self,
        *,
        name: str,
        description: str | None,
        storage_uri: str,
        schema_json: dict,
        row_count: int,
    ) -> TableAsset:
        return TableAsset(
            name=name,
            description=description,
            source_type=SourceType.UPLOAD_FILE,
            storage_type=StorageType.LOCAL_CSV,
            storage_uri=storage_uri,
            schema_json=schema_json,
            row_count=row_count,
        )

    def _normalize_xlsx_row(self, row: tuple[object, ...]) -> list[object]:
        return ["" if value is None else value for value in row]
